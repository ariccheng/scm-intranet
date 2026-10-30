"""Make the "Stores and menu list" spreadsheet for the Files app.

    uv run --with openpyxl python3 tools/make_stores_list.py

Run it from the scm-intranet folder after changing the stores or the menu
in content/company.json. It writes content/weeks/week-01/files/
stores-and-menu.xlsx with three sheets: Stores (code, name, year opened,
type), Menu (items, "New", ingredients, dietary tags) and Sauces. Public facts
only: no business data goes into this file.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "content" / "weeks" / "week-01" / "files" / "stores-and-menu.xlsx"


def sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"


def main():
    co = json.loads((ROOT / "content" / "company.json").read_text(encoding="utf-8"))
    menu = co.get("menu") or {}
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "Stores", ["Code", "Store", "Opened", "Type"],
          [[s.get("code"), s.get("name"), s.get("opened"), s.get("type")] for s in co.get("stores") or []], [8, 26, 10, 14])
    tag = {"vegetarian": "Vegetarian", "vegan": "Vegan", "gf": "Gluten-free", "spicy": "Spicy"}
    sheet(wb, "Menu", ["Item", "New", "What's in it", "Dietary"],
          [[m.get("name"), "New" if m.get("new") else "", m.get("description", ""), ", ".join(tag.get(x, x) for x in m.get("tags") or [])]
           for m in menu.get("items") or []], [32, 8, 70, 24])
    sauce = lambda x: [x, ""] if isinstance(x, str) else [x.get("name"), x.get("description", "")]
    sheet(wb, "Sauces", ["Sauce", "Description"], [sauce(x) for x in menu.get("sauces") or []], [28, 44])
    wb.properties.title = f'{co.get("name", "")} stores and menu'
    wb.properties.creator = co.get("name", "")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
